"""
parser.py -- Liest eine F2S-Workshop-Outcome-Excel und liefert normalisierte
Daten (Workshop-Metadaten + Eintragsliste) gemaess mapping.config.json.

Keine Hartkodierung von Spaltennamen/Statuswerten -- alles kommt aus der Config.
Unbekannte Status-/Typwerte werden NICHT verworfen, sondern als 'unbekannt'
normalisiert und als Warnung gemeldet.
"""
from __future__ import annotations

import json
from pathlib import Path
from datetime import datetime, date

import openpyxl


def load_config(config_path: str | Path) -> dict:
    with open(config_path, encoding="utf-8") as f:
        return json.load(f)


def _col_to_idx(letter: str) -> int:
    """Spaltenbuchstabe 'A' -> 1, 'B' -> 2, ... (openpyxl ist 1-basiert)."""
    idx = 0
    for ch in letter.strip().upper():
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    return idx


def _clean(value, placeholders) -> str:
    """Zelleninhalt zu getrimmtem String; Platzhalter werden zu ''."""
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    s = str(value).strip()
    if s in placeholders:
        return ""
    return s


def _norm_status(raw: str, status_map: dict, unknown: str):
    """Rohen Statuswert auf offen/in Arbeit/erledigt/... abbilden.
    Rueckgabe: (normalisiert, war_unbekannt)."""
    if not raw:
        return unknown, False  # leer ist kein 'unbekannter' Fehler
    if raw in status_map:
        return status_map[raw], False
    return unknown, True  # vorhanden, aber nicht in der Zuordnung -> melden


def parse_workbook(path: str | Path, config: dict) -> dict:
    """Parst eine Datei. Liefert:
    {
      "workshop": {...Metadaten...},
      "entries": [ {outcome_id, typ, lob, area, status, record_type, ...}, ... ],
      "warnings": [ {type, value, column, ...}, ... ]
    }
    """
    path = Path(path)
    placeholders = set(config.get("empty_placeholders", []))
    unknown = config.get("unknown_value", "unbekannt")
    warnings: list[dict] = []

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    # --- Workshop Summary -----------------------------------------------
    sc = config["summary_cells"]
    ws_sum = wb[config["sheets"]["summary"]]

    def cell(addr: str) -> str:
        return _clean(ws_sum[addr].value, placeholders)

    workshop = {
        "workshop_id": cell(sc["workshop_id"]),
        "workshop_title": cell(sc["workshop_title"]),
        "workstream": cell(sc["workstream"]),
        "area": cell(sc["area"]),
        "scope_items": cell(sc["scope_items"]),
        "date": cell(sc["date"]),
        "overall_result": cell(sc["overall_result"]),
        "approval_status": cell(sc["approval_status"]),
        "source_file": path.name,
    }
    # Workshop gilt als durchgefuehrt, wenn Overall Result gefuellt ist.
    workshop["conducted"] = bool(workshop["overall_result"])

    # Platzhalter-Erkennung: unausgefuellte Template-Kopien melden, nicht
    # stillschweigend als echten Workshop zaehlen.
    ph_markers = ("WS-ID", "e.g.", "[Enter", "[Name", "[List", "[Paste")
    def is_placeholder(v: str) -> bool:
        return any(m.lower() in v.lower() for m in ph_markers)

    if is_placeholder(workshop["workshop_id"]):
        warnings.append({"file": path.name, "type": "placeholder_workshop",
                         "value": workshop["workshop_id"], "column": "Workshop ID"})
        workshop["is_placeholder"] = True
        workshop["conducted"] = False  # unausgefuellte Vorlage zaehlt nicht

    lob = workshop["workstream"] or unknown
    area = workshop["area"] or unknown
    wsid = workshop["workshop_id"] or path.stem

    # --- Outcome Log ----------------------------------------------------
    log = wb[config["sheets"]["log"]]
    cols = {k: _col_to_idx(v) for k, v in config["log_columns"].items()}
    first_row = config["log_first_data_row"]

    rt_cfg = config["record_types"]
    rt_lookup = {v.lower(): k for k, v in rt_cfg.items()}  # "requirement"->"requirement"
    smaps = config["status_maps"]
    count_rules = config["count_rules"]

    entries: list[dict] = []

    def getval(row, key):
        c = cols.get(key)
        if not c:
            return ""
        return _clean(row[c - 1].value if c - 1 < len(row) else None, placeholders)

    for row in log.iter_rows(min_row=first_row):
        outcome_id = getval(row, "outcome_id")
        record_type_raw = getval(row, "record_type")
        if not outcome_id and not record_type_raw:
            continue  # Leerzeile

        # Zaehlmodell Variante B: Eine Outcome-Log-Zeile = GENAU EIN Eintrag,
        # klassifiziert nach Record Type. Eindeutig ueber (Datei + Outcome-ID).
        rt_key = rt_lookup.get(record_type_raw.lower()) if record_type_raw else None
        if record_type_raw and rt_key is None:
            warnings.append({"file": wsid, "type": "unknown_record_type",
                             "value": record_type_raw, "column": "Record Type"})
        typ = rt_key if rt_key in ("requirement", "issue", "assumption", "information") else unknown

        # Status der Zeile aus Fit Assessment (fuer Requirement/Issue massgeblich).
        fit_raw = getval(row, "fit_assessment")
        status, unk = _norm_status(fit_raw, smaps["fit_assessment"], unknown)
        if unk:
            warnings.append({"file": wsid, "type": "unknown_status",
                             "value": fit_raw, "column": "Fit Assessment"})

        # Task/Decision der Zeile als ATTRIBUTE (nicht als eigene Zaehlobjekte).
        decision_field = getval(row, count_rules["decision_if_field"])
        has_decision = bool(decision_field)
        decision_status = ""
        if has_decision:
            ds_raw = getval(row, "decision_status")
            decision_status, dunk = _norm_status(ds_raw, smaps["decision_status"], unknown)
            if dunk:
                warnings.append({"file": wsid, "type": "unknown_status",
                                 "value": ds_raw, "column": "Decision Status"})

        action_field = getval(row, count_rules["task_if_field"])
        has_task = bool(action_field)
        task_status = ""
        if has_task:
            as_raw = getval(row, "action_status")
            task_status, aunk = _norm_status(as_raw, smaps["action_status"], unknown)
            if aunk:
                warnings.append({"file": wsid, "type": "unknown_status",
                                 "value": as_raw, "column": "Action Status"})

        entries.append({
            # eindeutige Identitaet: Datei + Outcome-ID
            "uid": f"{path.name}|{outcome_id or '?'}",
            "outcome_id": outcome_id or f"{wsid}-?",
            "lob": lob,
            "area": area,
            "workshop_id": wsid,
            "workshop_title": workshop["workshop_title"],
            "workshop_date": workshop["date"],
            "typ": typ,
            "record_type": record_type_raw or unknown,
            "status": status,
            "fit_assessment": fit_raw or "",
            "process_step": getval(row, "process_step"),
            "topic": getval(row, "topic"),
            "business_requirement": getval(row, "business_requirement"),
            "criticality": getval(row, "criticality"),
            "proposed_solution": getval(row, "proposed_solution"),
            "owner": getval(row, "action_owner") or getval(row, "decision_owner"),
            "due_date": getval(row, "due_date"),
            "requirement_url": getval(row, "requirement_url"),
            # Attribute statt eigener Zaehlung:
            "has_task": has_task,
            "task_status": task_status,
            "has_decision": has_decision,
            "decision_status": decision_status,
            # Quell-Link (wird spaeter vom Downloader/refresh gesetzt):
            "source_file": path.name,
        })

    wb.close()
    return {"workshop": workshop, "entries": entries, "warnings": warnings}


if __name__ == "__main__":
    # Kleiner Selbsttest mit einer Datei als Argument.
    import sys
    cfg = load_config(Path(__file__).parent / "mapping.config.json")
    result = parse_workbook(sys.argv[1], cfg)
    print(json.dumps(result, indent=2, ensure_ascii=False))
