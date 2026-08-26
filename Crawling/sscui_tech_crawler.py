#!/usr/bin/env python3
"""
SSCUI Technical Information Crawler
Liest SSCUI-URLs aus einer Excel-Datei und crawlt die technischen
Feldinformationen via F1-Hilfe → Technical Information Popup.

Architektur: Alle DOM-Operationen laufen im Hauptseiten-Kontext mit
contentDocument-Zugriff auf den SAP-WebGUI-iFrame. Das ist nötig weil
die iFrames keine src-Attribute haben (dynamisch per JS befüllt) und
Playwright die Frames daher nicht direkt adressieren kann.
"""

import asyncio
import argparse
import base64
import io
import json
import os
import re
from pathlib import Path
from datetime import datetime

import anthropic
import pandas as pd
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from PIL import Image
from playwright.async_api import async_playwright, Page

OUTPUT_FILE = "sscui_technical_info.xlsx"
SHEET_NAME  = "Technische Daten"

TECH_COLUMNS = [
    "Screen - Program Name",
    "Screen - Screen Number",
    "GUI - Program Name",
    "GUI - Status",
    "Field - Table Name",
    "Field - Field Name",
    "Field - Data Element",
    "Batch - Screen Field",
]

ALL_COLUMNS = ["SSCUI Name", "SSCUI Nummer", "Screen", "Field Label"] + TECH_COLUMNS


# ---------------------------------------------------------------------------
# Login
# ---------------------------------------------------------------------------

async def manual_login_wait(page: Page):
    print("\n" + "=" * 60)
    print("MANUELLER LOGIN")
    print("1. Bitte im Browser in S/4HANA einloggen")
    print("2. Warte bis die SSCUI-Seite vollständig geladen ist")
    print("3. Drücke dann hier ENTER")
    print("=" * 60)
    input()


# ---------------------------------------------------------------------------
# iFrame-Index finden (läuft im Hauptseiten-Kontext)
# ---------------------------------------------------------------------------

async def find_content_iframe_index(page: Page) -> tuple:
    """
    Findet den iFrame-Index mit den meisten sichtbaren Input-Feldern.
    Läuft im Hauptseiten-JS-Kontext und greift via contentDocument auf
    die iFrames zu — genau wie im Browser DevTools Console möglich.
    Gibt (index, count) zurück. index=-1 bedeutet Hauptdokument.
    """
    await asyncio.sleep(3)  # Alle Frames laden lassen

    result = await page.evaluate("""
        () => {
            const INPUT_SEL = 'input:not([type="hidden"]):not([disabled]), select:not([disabled])';
            const iframes   = document.querySelectorAll('iframe');

            let bestIdx   = -1;
            let bestCount = 0;  // iframes always preferred over main document

            iframes.forEach((iframe, i) => {
                try {
                    const doc = iframe.contentDocument;
                    if (!doc || !doc.body) return;
                    const n = doc.querySelectorAll(INPUT_SEL).length;
                    if (n > bestCount) {
                        bestCount = n;
                        bestIdx   = i;
                    }
                } catch(e) {}
            });

            // Fallback to main document only when no iframe has inputs
            if (bestIdx === -1) {
                bestCount = document.querySelectorAll(INPUT_SEL).length;
            }

            return { idx: bestIdx, count: bestCount };
        }
    """)

    idx   = result['idx']
    count = result['count']

    if count > 0:
        label = f"iframe[{idx}]" if idx >= 0 else "Hauptdokument"
        print(f"  Content-Frame: {label}  ({count} Inputs gefunden)")
    else:
        # Diagnose: zeige was in jedem iFrame steckt
        diag = await page.evaluate("""
            () => {
                return Array.from(document.querySelectorAll('iframe')).map((f, i) => {
                    try {
                        const doc = f.contentDocument;
                        return {
                            idx: i,
                            inputs: doc ? doc.querySelectorAll('input').length : -1,
                            bodyLen: doc && doc.body ? doc.body.innerHTML.length : 0,
                        };
                    } catch(e) {
                        return { idx: i, inputs: -2, error: e.toString() };
                    }
                });
            }
        """)
        print("  WARNUNG: Keine Inputs gefunden. iFrame-Diagnose:")
        for d in diag:
            print(f"    iframe[{d['idx']}]: inputs={d.get('inputs')}, "
                  f"bodyLen={d.get('bodyLen', 0)}, err={d.get('error', '')[:60]}")

    return idx, count


# ---------------------------------------------------------------------------
# Seitentitel
# ---------------------------------------------------------------------------

async def get_page_title(page: Page, iframe_idx: int) -> str:
    """Liest den Seitentitel aus dem Content-iFrame."""
    return await page.evaluate("""
        (iframeIdx) => {
            const doc = iframeIdx >= 0
                ? document.querySelectorAll('iframe')[iframeIdx]?.contentDocument
                : document;
            if (!doc) return 'Unknown Screen';

            for (const sel of ['h1', '.sapMTitle', '.urSHdr', '[class*="Title"]']) {
                const el = doc.querySelector(sel);
                if (el) {
                    const t = el.innerText?.trim() || el.textContent?.trim() || '';
                    if (t.length > 3 && t.length < 200) return t;
                }
            }
            return doc.title || 'Unknown Screen';
        }
    """, iframe_idx)


# ---------------------------------------------------------------------------
# Claude Vision API – Felder per Screenshot identifizieren
# ---------------------------------------------------------------------------

def _load_api_key() -> str:
    """Lädt ANTHROPIC_API_KEY aus .env (Crawler-Verzeichnis oder SAP_Config_Bot)."""
    script_dir = Path(__file__).parent
    for search_path in [script_dir, script_dir.parent / "SAP_Config_Bot"]:
        env_file = search_path / ".env"
        if env_file.exists():
            load_dotenv(env_file)
            break
    else:
        load_dotenv()
    return os.getenv("ANTHROPIC_API_KEY", "")


_ANTHROPIC_CLIENT: anthropic.Anthropic | None = None


def _get_anthropic_client() -> anthropic.Anthropic:
    global _ANTHROPIC_CLIENT
    if _ANTHROPIC_CLIENT is None:
        key = _load_api_key()
        if not key:
            raise ValueError(
                "ANTHROPIC_API_KEY nicht gefunden. "
                "Erstelle .env im Crawler-Verzeichnis mit: ANTHROPIC_API_KEY=sk-ant-..."
            )
        _ANTHROPIC_CLIENT = anthropic.Anthropic(api_key=key)
    return _ANTHROPIC_CLIENT


def _compress_screenshot(screenshot_bytes: bytes, max_width: int = 1280, quality: int = 85):
    """Verkleinert und komprimiert Screenshot. Gibt (bytes, media_type, width, height) zurück."""
    try:
        img = Image.open(io.BytesIO(screenshot_bytes))
        if img.width > max_width:
            ratio = max_width / img.width
            img = img.resize((max_width, int(img.height * ratio)), Image.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True)
        return buf.getvalue(), "image/jpeg", img.width, img.height
    except Exception as e:
        print(f"  [Vision] Bild-Komprimierung fehlgeschlagen: {e}")
        return screenshot_bytes, "image/png", 1280, 900


async def identify_fields_via_screenshot(page: Page) -> list:
    """
    Macht einen Seiten-Screenshot, schickt ihn an Claude Vision.
    Claude identifiziert alle sichtbaren Tabellenspalten / Formularfelder
    und liefert Klick-Koordinaten der ersten Datenzeile zurück.
    """
    screenshot_bytes = await page.screenshot()
    viewport = page.viewport_size or {"width": 1280, "height": 900}
    vp_w, vp_h = viewport["width"], viewport["height"]

    compressed, media_type, comp_w, comp_h = _compress_screenshot(screenshot_bytes)
    b64 = base64.standard_b64encode(compressed).decode("utf-8")

    # Skalierungsfaktoren: komprimierter Screenshot → Viewport CSS-Pixel
    scale_x = vp_w / comp_w
    scale_y = vp_h / comp_h

    prompt = (
        "This is a screenshot of a SAP S/4HANA SSCUI maintenance screen.\n"
        "Identify ALL visible table column headers. For each column give the approximate pixel "
        "coordinates (x, y) of a clickable cell in the FIRST DATA ROW "
        "(the row directly below the headers — not the header row itself).\n\n"
        "If the screen shows a form (not a table), list all visible form field labels "
        "with the coordinates of their input elements.\n\n"
        f"The screenshot is {comp_w}x{comp_h} pixels. Report coordinates in screenshot pixel space.\n\n"
        "Return ONLY a valid JSON array, no markdown, no explanation:\n"
        '[{"label": "CCtC", "x": 150, "y": 320}, {"label": "Name", "x": 280, "y": 320}, ...]'
    )

    try:
        client = _get_anthropic_client()
        response = client.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=1000,
            messages=[{
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {"type": "base64", "media_type": media_type, "data": b64},
                    },
                    {"type": "text", "text": prompt},
                ],
            }],
        )
    except Exception as e:
        print(f"  [Vision] API-Fehler: {e}")
        return []

    raw = response.content[0].text.strip()
    raw = re.sub(r"^```(?:json)?", "", raw, flags=re.MULTILINE).strip()
    raw = re.sub(r"```$",          "", raw, flags=re.MULTILINE).strip()

    try:
        fields_raw = json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"  [Vision] JSON-Parsing fehlgeschlagen: {e}")
        print(f"  [Vision] Antwort: {raw[:300]}")
        return []

    # SAP-Zeilenauswahl-Spalten überspringen (nicht crawlen)
    SKIP_LABELS = {'select', 'sel', 'sel.', 'selection', 'auswahl', 'row', 'zeile', 'z.'}

    fields = []
    seen: set = set()
    for f in fields_raw:
        label = str(f.get("label", "")).strip()
        if not label or label in seen:
            continue
        if label.lower() in SKIP_LABELS:
            continue
        seen.add(label)
        x = float(f.get("x", 0)) * scale_x
        y = float(f.get("y", 0)) * scale_y
        fields.append({"label": label, "x": x, "y": y, "type": "col"})

    print(f"  [Vision] {len(fields)} Felder erkannt: {[f['label'] for f in fields]}")
    return fields


# ---------------------------------------------------------------------------
# Felder finden (main-page-relative Koordinaten) – DOM-Fallback
# ---------------------------------------------------------------------------

async def find_visible_fields(page: Page, iframe_idx: int) -> list:
    """
    Strategie 1 – Tabellen-Ansicht (SM30/Pflegeansicht):
      Findet Header-Zeile (erste Zeile >= 2 beschriftete Zellen, kein Input)
      und aktive Datenzeile (erste Zeile mit Inputs). Klickt alle Datenzellen –
      deckt Text-Felder, Checkboxen und Read-Only-Felder ab.
    Strategie 2 – Koordinaten-basiert (separate Header/Daten-Tabellen):
      Sucht Textelemente direkt über der Datenzeile (SAP-Split-Tabellen-Layout).
    Strategie 3 – Formular-Ansicht: sichtbare Inputs mit Adjacent-Label.
    """
    return await page.evaluate(r"""
        (iframeIdx) => {
            var iframe    = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
            var doc       = iframe ? iframe.contentDocument : document;
            var iframeBox = iframe ? iframe.getBoundingClientRect() : { left: 0, top: 0 };
            if (!doc || !doc.body) return [];

            var fields     = [];
            var seenLabels = {};
            var viewH      = doc.defaultView ? doc.defaultView.innerHeight : 800;

            function cleanLabel(txt) {
                return (txt || '').replace(/\s+/g, ' ').trim();
            }

            function addField(label, rect, type) {
                if (!label || seenLabels[label]) return;
                seenLabels[label] = true;
                fields.push({
                    label: label, type: type || 'col',
                    x: iframeBox.left + rect.left + rect.width  / 2,
                    y: iframeBox.top  + rect.top  + rect.height / 2
                });
            }

            // ── Strategie 1: gleiche Tabelle hat Header-Zeile + aktive Daten-Zeile ──
            var tables = Array.from(doc.querySelectorAll('table'));
            for (var ti = 0; ti < tables.length; ti++) {
                var allRows = Array.from(tables[ti].querySelectorAll('tr'));
                if (allRows.length < 2) continue;

                var hdrRow  = null;   // erste Zeile >= 2 beschriftete Zellen, kein Input
                var dataRow = null;   // erste Zeile MIT Inputs (= aktive Zeile)

                for (var ri = 0; ri < allRows.length; ri++) {
                    var row = allRows[ri];
                    if (row.querySelector('input, select, textarea')) {
                        if (hdrRow) { dataRow = row; break; }
                    } else if (!hdrRow) {
                        // Nur EINMAL setzen – erste gültige Zeile = Header
                        var lc = Array.from(row.children).filter(function(c) {
                            return cleanLabel(c.textContent).length > 0;
                        }).length;
                        if (lc >= 2) hdrRow = row;
                    }
                }
                if (!hdrRow || !dataRow) continue;

                var hdrCells  = Array.from(hdrRow.children);
                var dataCells = Array.from(dataRow.children);
                var found = 0;
                for (var ci = 0; ci < hdrCells.length; ci++) {
                    var lbl = cleanLabel(hdrCells[ci].textContent);
                    if (!lbl) continue;
                    var cell = dataCells[ci];
                    if (!cell) continue;
                    var r = cell.getBoundingClientRect();
                    if (r.width === 0 || r.height === 0) continue;
                    if (r.top > viewH + 100) continue;
                    addField(lbl, r, 'col');
                    found++;
                }
                if (found > 0) break;
            }

            // ── Strategie 2: Header und Daten in GETRENNTEN Tabellen (SAP-Split) ──
            // Findet aktive Datenzeile und sucht Textelemente direkt darüber.
            if (fields.length === 0) {
                var activeRow = null;
                for (var ti2 = 0; ti2 < tables.length; ti2++) {
                    var tr = tables[ti2].querySelector('tr');
                    while (tr) {
                        if (tr.querySelector('input:not([type="hidden"]), select')) {
                            activeRow = tr; break;
                        }
                        tr = tr.nextElementSibling;
                    }
                    if (activeRow) break;
                }

                if (activeRow) {
                    var aRect = activeRow.getBoundingClientRect();
                    Array.from(activeRow.children).forEach(function(cell) {
                        var cr = cell.getBoundingClientRect();
                        if (cr.width === 0 || cr.height === 0) return;
                        var cx = cr.left + cr.width / 2;
                        // Nächstes Textelement über der aktiven Zeile an dieser x-Position
                        var best = null; var bestD = 9999;
                        Array.from(doc.querySelectorAll('td, th, span, div')).forEach(function(el) {
                            if (el.querySelector('input, select, textarea')) return;
                            var er = el.getBoundingClientRect();
                            if (er.height === 0 || er.width === 0) return;
                            if (er.bottom >= aRect.top) return;      // muss oberhalb sein
                            if (er.top < aRect.top - 80) return;     // max 80px Abstand
                            if (er.right < cx || er.left > cx) return;
                            var txt = cleanLabel(el.textContent);
                            if (!txt || txt.length > 40) return;
                            var d = aRect.top - er.bottom;
                            if (d < bestD) { bestD = d; best = el; }
                        });
                        if (best) addField(cleanLabel(best.textContent), cr, 'col');
                    });
                }
            }

            // ── Strategie 3: Formular-Ansicht – sichtbare Inputs ─────────────────
            if (fields.length === 0) {
                var INPUT_SEL = 'input:not([type="hidden"]):not([disabled]), select:not([disabled])';
                Array.from(doc.querySelectorAll(INPUT_SEL)).forEach(function(el) {
                    var r = el.getBoundingClientRect();
                    if (r.width === 0 || r.height === 0) return;
                    if (r.top > viewH + 100) return;
                    var node = el;
                    while (node && node !== doc.body) {
                        var s = (doc.defaultView || window).getComputedStyle(node);
                        if (s.display === 'none' || s.visibility === 'hidden') return;
                        node = node.parentElement;
                    }
                    var label = '';
                    var pTd = el.closest ? el.closest('td') : null;
                    if (pTd && pTd.previousElementSibling)
                        label = cleanLabel(pTd.previousElementSibling.textContent);
                    if (!label) {
                        var aria = el.getAttribute('aria-label');
                        if (aria && aria.trim() !== (el.value || '').trim()) label = aria.trim();
                    }
                    if (!label) {
                        var ttl = el.getAttribute('title');
                        if (ttl && ttl.trim() !== (el.value || '').trim()) label = ttl.trim();
                    }
                    if (!label && el.id) {
                        try {
                            var lb = doc.querySelector('label[for="' + el.id + '"]');
                            if (lb) label = cleanLabel(lb.textContent);
                        } catch(e) {}
                    }
                    if (!label) label = 'Feld_' + (fields.length + 1);
                    addField(label, r, el.type || 'input');
                });
            }

            fields.sort(function(a, b) {
                var rA = Math.round(a.y / 15), rB = Math.round(b.y / 15);
                return rA !== rB ? rA - rB : a.x - b.x;
            });
            return fields;
        }
    """, iframe_idx)


# ---------------------------------------------------------------------------
# Popup-Erkennung & Button-Klick im iFrame
# ---------------------------------------------------------------------------

POPUP_SEL = '[role="dialog"], .urPopup, .urPopupWindow, .urPopupBorder, .sapMDialog, .sapUiDlg'

async def wait_for_popup(page: Page, iframe_idx: int, timeout_s: float = 3.5) -> bool:
    """Wartet bis ein Popup im iFrame sichtbar ist."""
    import time
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        found = await page.evaluate("""
            ([iframeIdx, popupSel]) => {
                const iframe = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
                const doc    = iframe ? iframe.contentDocument : document;
                if (!doc) return false;
                const popups = [...doc.querySelectorAll(popupSel)];
                return popups.some(p => {
                    const s = (doc.defaultView || window).getComputedStyle(p);
                    return s.display !== 'none' && s.visibility !== 'hidden';
                });
            }
        """, [iframe_idx, POPUP_SEL])
        if found:
            return True
        await asyncio.sleep(0.3)
    return False


async def click_button_in_iframe(page: Page, iframe_idx: int, *button_texts) -> bool:
    """
    Findet einen Button/Link im iFrame anhand seines Textes und klickt ihn
    über main-page-relative Koordinaten.
    """
    for text in button_texts:
        coords = await page.evaluate("""
            ([iframeIdx, btnText]) => {
                const iframe    = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
                const doc       = iframe ? iframe.contentDocument : document;
                const iframeBox = iframe ? iframe.getBoundingClientRect() : { left: 0, top: 0 };
                if (!doc) return null;

                const candidates = doc.querySelectorAll('a, button, span, td, [onclick]');
                for (const el of candidates) {
                    if (el.innerText?.trim() === btnText || el.textContent?.trim() === btnText) {
                        const rect = el.getBoundingClientRect();
                        if (rect.width === 0 || rect.height === 0) continue;
                        return {
                            x: iframeBox.left + rect.left + rect.width  / 2,
                            y: iframeBox.top  + rect.top  + rect.height / 2,
                        };
                    }
                }
                return null;
            }
        """, [iframe_idx, text])

        if coords:
            await page.mouse.click(coords['x'], coords['y'])
            return True
    return False


async def click_tech_info_in_popup(page: Page, iframe_idx: int) -> bool:
    """
    Klickt 'Technical Information' / 'Technische Informationen' NUR innerhalb
    des sichtbaren Popups – verhindert versehentliches Klicken auf gleichnamige
    Elemente außerhalb des Dialogs.
    """
    coords = await page.evaluate(r"""
        ([iframeIdx, popupSel]) => {
            var iframe    = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
            var doc       = iframe ? iframe.contentDocument : document;
            var iframeBox = iframe ? iframe.getBoundingClientRect() : {left:0, top:0};
            if (!doc) return null;

            var popups = Array.from(doc.querySelectorAll(popupSel)).filter(function(p) {
                try {
                    var s = (doc.defaultView || window).getComputedStyle(p);
                    return s.display !== 'none' && s.visibility !== 'hidden';
                } catch(e) { return false; }
            });

            var btnTexts = ['Technical Information', 'Technische Informationen'];
            for (var pi = popups.length - 1; pi >= 0; pi--) {
                var popup = popups[pi];
                var cands = popup.querySelectorAll('a, button, span, td, [onclick]');
                for (var ci = 0; ci < cands.length; ci++) {
                    var txt = (cands[ci].innerText || cands[ci].textContent || '').trim();
                    if (btnTexts.indexOf(txt) >= 0) {
                        var r = cands[ci].getBoundingClientRect();
                        if (r.width > 0 && r.height > 0)
                            return {x: iframeBox.left + r.left + r.width/2,
                                    y: iframeBox.top  + r.top  + r.height/2};
                    }
                }
            }
            return null;
        }
    """, [iframe_idx, POPUP_SEL])

    if coords:
        await page.mouse.click(coords['x'], coords['y'])
        return True
    return False


async def close_all_popups(page: Page, iframe_idx: int):
    """Schließt alle offenen Popups über den roten X-Button oder Escape."""
    for _ in range(4):
        coords = await page.evaluate(r"""
            ([iframeIdx, popupSel]) => {
                var iframe    = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
                var doc       = iframe ? iframe.contentDocument : document;
                var iframeBox = iframe ? iframe.getBoundingClientRect() : {left:0, top:0};
                if (!doc) return null;

                var popups = Array.from(doc.querySelectorAll(popupSel)).filter(function(p) {
                    try {
                        var s = (doc.defaultView || window).getComputedStyle(p);
                        return s.display !== 'none' && s.visibility !== 'hidden';
                    } catch(e) { return false; }
                });
                if (popups.length === 0) return null;

                var popup = popups[popups.length - 1];

                // 1. aria-label / title / class
                var closeSelectors = [
                    '[aria-label="Close"]', '[aria-label="Schließen"]',
                    '[title="Close"]',      '[title="Schließen"]',
                    '.urBtnClose', '.urHdlClose',
                    'button[class*="lose"]', 'a[class*="lose"]'
                ];
                for (var si = 0; si < closeSelectors.length; si++) {
                    try {
                        var btn = popup.querySelector(closeSelectors[si]);
                        if (btn) {
                            var r = btn.getBoundingClientRect();
                            if (r.width > 0 && r.height > 0)
                                return {x: iframeBox.left + r.left + r.width/2,
                                        y: iframeBox.top  + r.top  + r.height/2};
                        }
                    } catch(e) {}
                }

                // 2. × / ✕ / X Symbol
                var allEls = popup.querySelectorAll('a, button, span, div');
                for (var ei = 0; ei < allEls.length; ei++) {
                    var txt = (allEls[ei].innerText || allEls[ei].textContent || '').trim();
                    if (txt === '×' || txt === '✕' || txt === 'X') {
                        var r2 = allEls[ei].getBoundingClientRect();
                        if (r2.width > 0 && r2.height > 0)
                            return {x: iframeBox.left + r2.left + r2.width/2,
                                    y: iframeBox.top  + r2.top  + r2.height/2};
                    }
                }
                return null;
            }
        """, [iframe_idx, POPUP_SEL])

        if coords:
            await page.mouse.click(coords['x'], coords['y'])
            await asyncio.sleep(0.3)
        else:
            try:
                await page.keyboard.press('Escape')
                await asyncio.sleep(0.2)
            except Exception:
                break

        # Prüfen ob noch Popups offen
        still_open = await page.evaluate("""
            ([iframeIdx, popupSel]) => {
                var iframe = iframeIdx >= 0 ? document.querySelectorAll('iframe')[iframeIdx] : null;
                var doc    = iframe ? iframe.contentDocument : document;
                if (!doc) return false;
                return Array.from(doc.querySelectorAll(popupSel)).some(function(p) {
                    try {
                        var s = (doc.defaultView || window).getComputedStyle(p);
                        return s.display !== 'none' && s.visibility !== 'hidden';
                    } catch(e) { return false; }
                });
            }
        """, [iframe_idx, POPUP_SEL])
        if not still_open:
            break


# ---------------------------------------------------------------------------
# Technical Information extrahieren
# ---------------------------------------------------------------------------

async def extract_tech_info(page: Page, iframe_idx: int) -> dict:
    """
    Extrahiert alle Felder aus dem Technical Information Popup.
    SAP WebDynpro rendert Werte in <input readonly>-Feldern.
    Sucht in ALLEN iFrames + Hauptdokument.
    WICHTIG: raw string r-prefix damit Python \n nicht als Newline interpretiert.
    """
    raw = await page.evaluate(r"""
        ([popupSel]) => {
            // Alle Dokumente sammeln: Hauptseite + alle iFrames
            var allDocs = [{ doc: document, label: 'main' }];
            Array.from(document.querySelectorAll('iframe')).forEach(function(f, i) {
                try {
                    var d = f.contentDocument;
                    if (d) allDocs.push({ doc: d, label: 'iframe[' + i + ']' });
                } catch(e) {}
            });

            var TARGET = ['Screen Number','Data Element','Table Name',
                          'Datenelement','Tabellenname','Field Name','Feldname',
                          'Program Name','Programm','Screen Data','Bildschirmdaten',
                          'GUI Data','GUI-Daten','Field Description'];

            var dialog  = null;
            var foundIn = '';

            // Strategie 1: CSS-Selektor
            for (var di = 0; di < allDocs.length; di++) {
                var doc   = allDocs[di].doc;
                var lbl   = allDocs[di].label;
                var vis   = Array.from(doc.querySelectorAll(popupSel)).filter(function(d) {
                    try {
                        var s = (d.ownerDocument && d.ownerDocument.defaultView
                                 ? d.ownerDocument.defaultView : window).getComputedStyle(d);
                        return s.display !== 'none' && s.visibility !== 'hidden';
                    } catch(e) { return false; }
                });
                for (var vi = vis.length - 1; vi >= 0; vi--) {
                    var cand = vis[vi];
                    if (TARGET.some(function(t) { return cand.textContent.includes(t); })) {
                        dialog = cand; foundIn = lbl; break;
                    }
                }
                if (dialog) break;
            }

            // Strategie 2: beliebiges sichtbares Element mit Target-Text und Inputs
            if (!dialog) {
                for (var di2 = 0; di2 < allDocs.length; di2++) {
                    var doc2 = allDocs[di2].doc;
                    var lbl2 = allDocs[di2].label;
                    var els  = Array.from(doc2.querySelectorAll('table, div[class], form'));
                    for (var ei = 0; ei < els.length; ei++) {
                        var el = els[ei];
                        if (!TARGET.some(function(t) { return el.textContent.includes(t); })) continue;
                        if (el.querySelectorAll('input').length === 0) continue;
                        try {
                            var es = (doc2.defaultView || window).getComputedStyle(el);
                            if (es.display === 'none' || es.visibility === 'hidden') continue;
                        } catch(e) { continue; }
                        dialog = el; foundIn = lbl2 + '-fb'; break;
                    }
                    if (dialog) break;
                }
            }

            if (!dialog) {
                var diag = allDocs.map(function(entry) {
                    var d = entry.doc;
                    return entry.label + ':' +
                           d.querySelectorAll(popupSel).length + 'pop/' +
                           d.querySelectorAll('input').length + 'inp/' +
                           (TARGET.some(function(t) {
                               return d.body && d.body.textContent.includes(t);
                           }) ? 'HAS' : 'no');
                }).join('|');
                return { _status: 'no_dialog', _diag: diag };
            }

            var sectionKeyMap = {
                'Screen Data': {
                    'Program Name':  'Screen - Program Name',
                    'Screen Number': 'Screen - Screen Number'
                },
                'GUI Data': {
                    'Program Name': 'GUI - Program Name',
                    'Status':       'GUI - Status'
                },
                'Field Data': {
                    'Table Name':   'Field - Table Name',
                    'Field Name':   'Field - Field Name',
                    'Data Element': 'Field - Data Element'
                },
                'Field Description for Batch Input': {
                    'Screen Field': 'Batch - Screen Field'
                },
                'Bildschirmdaten': {
                    'Programm': 'Screen - Program Name',
                    'Dynpro':   'Screen - Screen Number'
                },
                'GUI-Daten': {
                    'Programm': 'GUI - Program Name',
                    'Status':   'GUI - Status'
                },
                'Felddaten': {
                    'Tabellenname': 'Field - Table Name',
                    'Feldname':     'Field - Field Name',
                    'Datenelement': 'Field - Data Element'
                },
                'Feldbeschreibung fur Batch-Input': {
                    'Bildschirmfeld': 'Batch - Screen Field'
                }
            };
            var sectionHdrs = Object.keys(sectionKeyMap);

            // SAP rendert Labels mit erstem Zeichen in separatem Span:
            // "P rogram Name" → "Program Name"
            function normLbl(raw) {
                return raw.trim().replace(/^([A-Z]) ([a-z])/, '$1$2');
            }

            function cellVal(cell) {
                var inp = cell.querySelector('input:not([type="hidden"])');
                if (inp) return (inp.value || inp.getAttribute('value') || '').trim();
                var sel = cell.querySelector('select');
                if (sel && sel.options.length > 0 && sel.selectedIndex >= 0)
                    return (sel.options[sel.selectedIndex].text || '').trim();
                return (cell.textContent || '').trim();
            }

            // Das Popup hat IMMER dieselbe Struktur mit denselben Inputs in fixer Reihenfolge.
            // Daher: Inputs nach DOM-Position sammeln und direkt den Spalten zuordnen.
            // Das umgeht das Label-Split-Problem ("P rogram Name" statt "Program Name").
            var FIXED_COLS = [
                'Screen - Program Name',
                'Screen - Screen Number',
                'GUI - Program Name',
                'GUI - Status',
                'Field - Table Name',
                'Field - Field Name',
                'Field - Data Element',
                'Batch - Screen Field'
            ];

            var result   = {};
            var allInps  = Array.from(dialog.querySelectorAll('input:not([type="hidden"])'));
            allInps.forEach(function(inp, i) {
                if (i < FIXED_COLS.length) {
                    var v = (inp.value || inp.getAttribute('value') || '').trim();
                    if (v) result[FIXED_COLS[i]] = v;
                }
            });

            // Diagnostics: Text + input values
            var diagParts = [];
            function diagCollect(node) {
                if (node.nodeType === 3) {
                    var t2 = node.textContent.trim();
                    if (t2) diagParts.push(t2);
                } else if (node.tagName === 'INPUT' && node.type !== 'hidden') {
                    var v2 = (node.value || '').trim();
                    if (v2) diagParts.push('[' + v2 + ']');
                } else { node.childNodes.forEach(diagCollect); }
            }
            diagCollect(dialog);

            var dataCount = Object.keys(result).filter(function(k) { return k[0] !== '_'; }).length;
            result._status   = dataCount > 0 ? 'ok' : 'empty';
            result._foundIn  = foundIn;
            result._diagText = diagParts.join(' ').slice(0, 400);
            return result;
        }
    """, [POPUP_SEL])

    if not isinstance(raw, dict):
        return {}
    status    = raw.pop('_status',   '')
    found_in  = raw.pop('_foundIn',  '')
    diag      = raw.pop('_diag',     '')
    diag_text = raw.pop('_diagText', '')

    if status != 'ok':
        print(f"\n    [popup] status={status!r}  foundIn={found_in!r}")
        if diag:
            print(f"    [popup diag] {diag}")
        if diag_text:
            print(f"    [popup text] {repr(diag_text[:300])}")
    return raw


# ---------------------------------------------------------------------------
# SSCUI-Seite verarbeiten
# ---------------------------------------------------------------------------

async def process_sscui_page(page: Page, sscui_name: str, sscui_nr: str) -> list:
    results = []

    iframe_idx, field_count = await find_content_iframe_index(page)

    screen_title = await get_page_title(page, iframe_idx)
    print(f"  Screen: {screen_title}")

    if field_count == 0:
        print("  Keine Felder – überspringe.")
        return results

    fields = await identify_fields_via_screenshot(page)
    if not fields:
        print("  [Vision] Kein Ergebnis – versuche DOM-Fallback...")
        fields = await find_visible_fields(page, iframe_idx)
    print(f"  {len(fields)} eindeutige Felder")

    if not fields:
        print("  Felder nicht sichtbar – überspringe.")
        return results

    for i, field in enumerate(fields):
        label = field['label']
        x, y  = field['x'], field['y']

        print(f"  [{i+1:>2}/{len(fields)}] '{label}' ({field['type']})...",
              end=" ", flush=True)

        tech_data = {}
        try:
            for attempt in range(2):
                if attempt == 0:
                    # Erstversuch: Feld anklicken
                    await page.mouse.click(x, y)
                    await asyncio.sleep(0.5)

                    # Falls durch Klick ein Value-Help-Popup geöffnet wurde
                    # (z.B. bei read-only Feldern wie CCtC) → vor F1 schließen
                    if await wait_for_popup(page, iframe_idx, timeout_s=0.8):
                        await page.keyboard.press('Escape')
                        await asyncio.sleep(0.3)
                else:
                    # Retry: Popup schließen, KEIN erneuter Klick –
                    # verhindert Spalten-Drift durch SAP-Neurendering
                    await close_all_popups(page, iframe_idx)
                    await asyncio.sleep(0.5)
                    print("(Retry)...", end=" ", flush=True)

                # F1 drücken
                await page.keyboard.press('F1')
                await asyncio.sleep(2.0)

                # Auf F1-Hilfe-Popup warten
                if not await wait_for_popup(page, iframe_idx):
                    if attempt == 0:
                        continue  # noch ein Versuch
                    print("kein F1-Popup")
                    break

                # "Technical Information" NUR innerhalb des Popups klicken
                clicked = await click_tech_info_in_popup(page, iframe_idx)
                if not clicked:
                    await close_all_popups(page, iframe_idx)
                    if attempt == 0:
                        continue  # noch ein Versuch
                    print("kein Tech-Info Button")
                    break

                await asyncio.sleep(2.5)  # Tech-Info-Popup vollständig laden lassen

                # Daten extrahieren
                tech_data = await extract_tech_info(page, iframe_idx)
                if tech_data:
                    break  # Erfolg – kein weiterer Versuch nötig
                # Kein Ergebnis → Popup schließen und nochmal versuchen
                if attempt == 1:
                    print("keine Daten im Popup")

            if tech_data:
                row = {
                    "SSCUI Name":   sscui_name,
                    "SSCUI Nummer": sscui_nr,
                    "Screen":       screen_title,
                    "Field Label":  label,
                }
                for col in TECH_COLUMNS:
                    row[col] = tech_data.get(col, '')
                results.append(row)
                print(f"OK  ({tech_data.get('Field - Field Name', '')})")

        except Exception as e:
            print(f"FEHLER: {e}")

        # Popup immer schließen bevor das nächste Feld angeklickt wird
        await close_all_popups(page, iframe_idx)
        await asyncio.sleep(0.3)

    return results


# ---------------------------------------------------------------------------
# Excel-Ausgabe
# ---------------------------------------------------------------------------

def write_excel(output_path: Path, all_results: list):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")

    for col, h in enumerate(ALL_COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_idx, row_data in enumerate(all_results, 2):
        for col_idx, col_name in enumerate(ALL_COLUMNS, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_name, ''))
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(vertical="top")

    col_widths = [40, 15, 45, 25, 22, 18, 22, 20, 20, 28]
    for col, w in enumerate(col_widths[:len(ALL_COLUMNS)], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w

    ws.freeze_panes = "A2"

    try:
        wb.save(output_path)
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = output_path.with_stem(output_path.stem + f"_{ts}")
        wb.save(output_path)
        print(f"  Datei gesperrt – gespeichert als: {output_path.name}")

    print(f"  {len(all_results)} Einträge gespeichert: {output_path.resolve()}")


# ---------------------------------------------------------------------------
# Haupt-Crawler
# ---------------------------------------------------------------------------

async def run_crawler(input_path: Path, output_path: Path, start_idx: int = 0,
                      save_every: int = 10):
    df = pd.read_excel(input_path)

    def find_col(keywords):
        for kw in keywords:
            for c in df.columns:
                if kw.lower() in str(c).lower():
                    return c
        return df.columns[0]

    name_col = find_col(['sscui name', 'name'])
    nr_col   = find_col(['nummer', 'number', 'nr', 'id'])
    url_col  = find_col(['link', 'url'])

    rows = df[[name_col, nr_col, url_col]].copy()
    rows = rows[rows[url_col].astype(str).str.startswith('http', na=False)].reset_index(drop=True)

    if start_idx > 0:
        rows = rows.iloc[start_idx:].reset_index(drop=True)

    print(f"Eingabe:  {input_path.name}")
    print(f"SSCUIs:   {len(rows)} (ab Index {start_idx})")
    print(f"Ausgabe:  {output_path}")

    all_results = []

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False)
        page    = await (await browser.new_context()).new_page()

        first_url = str(rows.iloc[0][url_col]).strip()
        print(f"\nNavigiere zu: {first_url[:80]}")
        try:
            await page.goto(first_url, wait_until="commit", timeout=30000)
        except Exception:
            pass

        await manual_login_wait(page)

        for idx, row_data in rows.iterrows():
            sscui_name = str(row_data[name_col]).strip()
            sscui_nr   = str(row_data[nr_col]).strip()
            url        = str(row_data[url_col]).strip()

            print(f"\n[{idx+1}/{len(rows)}] {sscui_name} ({sscui_nr})")

            try:
                if idx > 0:
                    await page.goto(url, wait_until="commit", timeout=30000)
                    try:
                        await page.wait_for_load_state("networkidle", timeout=20000)
                    except Exception:
                        pass
                    await asyncio.sleep(3)

                page_results = await process_sscui_page(page, sscui_name, sscui_nr)
                all_results.extend(page_results)
                print(f"  → {len(page_results)} Felder mit Tech-Info")

            except Exception as e:
                print(f"  FEHLER: {e}")

            if all_results and (idx + 1) % save_every == 0:
                print(f"  [Zwischenspeichern nach {idx+1} SSCUIs...]")
                write_excel(output_path, all_results)

        try:
            await browser.close()
        except Exception:
            pass

    if all_results:
        print(f"\n{len(all_results)} Einträge total – speichere...")
        write_excel(output_path, all_results)
    else:
        print("\nKeine Daten gecrawlt.")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="SSCUI Technical Info Crawler – F1-Technische Informationen",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Beispiele:
  python sscui_tech_crawler.py -i sscui_export.xlsx
  python sscui_tech_crawler.py -i sscui_export.xlsx -o ergebnis.xlsx
  python sscui_tech_crawler.py -i sscui_export.xlsx --start 50
        """,
    )
    parser.add_argument("--input",  "-i", required=True,
                        help="Excel-Datei mit SSCUI Name, Nummer und Link")
    parser.add_argument("--output", "-o", default=OUTPUT_FILE,
                        help=f"Ausgabe-Excel (Standard: {OUTPUT_FILE})")
    parser.add_argument("--start",  "-s", type=int, default=0,
                        help="Index ab dem gestartet wird (0-basiert)")
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"FEHLER: Datei nicht gefunden: {input_path}")
        return

    asyncio.run(run_crawler(input_path, Path(args.output), args.start))


if __name__ == "__main__":
    main()
